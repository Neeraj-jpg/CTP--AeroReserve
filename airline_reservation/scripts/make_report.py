"""Generate the fact-checked AeroReserve project report (DOCX)."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"
DIAGRAMS = DOCS / "diagrams"
OUT = DOCS / "PROJECT_REPORT.docx"
TEST_COUNT = 45
COVERAGE = "93.1%"
REPORT_DATE = "13 October 2026"

NAVY = "102B46"
BLUE = "1463A5"
CYAN = "14B8C5"
GOLD = "F4A62A"
INK = "172B3E"
MUTED = "536B80"
PALE = "F3F7FA"
LIGHT_BLUE = "E8F3F9"
WHITE = "FFFFFF"
GREEN = "16835D"


def _rgb(value: str) -> RGBColor:
    return RGBColor.from_string(value)


def _shade(cell, fill: str) -> None:
    props = cell._tc.get_or_add_tcPr()
    shd = props.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        props.append(shd)
    shd.set(qn("w:fill"), fill)


def _cell_margins(cell, top: int = 90, start: int = 110, bottom: int = 90, end: int = 110) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def _keep_row(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def _add_page_field(paragraph) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for item in (begin, instr, separate, text, end):
        run._r.append(item)
    run.font.name = "Calibri"
    run.font.size = Pt(9)
    run.font.color.rgb = _rgb(MUTED)


def _set_styles(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = _rgb(INK)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.12
    normal.paragraph_format.widow_control = True
    for name, size, color in (
        ("Title", 28, NAVY),
        ("Heading 1", 19, NAVY),
        ("Heading 2", 14, BLUE),
        ("Heading 3", 11.5, CYAN),
    ):
        style = doc.styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = _rgb(color)
        style.paragraph_format.space_before = Pt(12 if name != "Title" else 2)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.keep_with_next = True
    for name in ("List Bullet", "List Number"):
        style = doc.styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(10.5)
        style.paragraph_format.space_after = Pt(3)
        style.paragraph_format.left_indent = Inches(0.22)
        style.paragraph_format.first_line_indent = Inches(-0.16)
    code = doc.styles.add_style("Report Code", WD_STYLE_TYPE.PARAGRAPH)
    code.font.name = "Consolas"
    code.font.size = Pt(8.5)
    code.font.color.rgb = _rgb(INK)
    code.paragraph_format.left_indent = Inches(0.14)
    code.paragraph_format.right_indent = Inches(0.1)
    code.paragraph_format.space_before = Pt(3)
    code.paragraph_format.space_after = Pt(5)
    code.paragraph_format.line_spacing = 1.0


def _header_footer(section) -> None:
    section.header_distance = Inches(0.42)
    section.footer_distance = Inches(0.42)
    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = header.add_run("AERORESERVE  |  PROJECT DEFENCE  |  OCTOBER 2026")
    run.font.name = "Calibri"
    run.font.size = Pt(8)
    run.font.bold = True
    run.font.color.rgb = _rgb(MUTED)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    lead = footer.add_run("AeroReserve · Academic demo · Page ")
    lead.font.name = "Calibri"
    lead.font.size = Pt(8)
    lead.font.color.rgb = _rgb(MUTED)
    _add_page_field(footer)


def _paragraph(
    doc: Document,
    text: str = "",
    *,
    style: str | None = None,
    align=None,
    italic=False,
    color=None,
    size=None,
    bold=False,
):
    paragraph = doc.add_paragraph(style=style)
    if align is not None:
        paragraph.alignment = align
    run = paragraph.add_run(text)
    run.italic = italic
    run.bold = bold
    if color:
        run.font.color.rgb = _rgb(color)
    if size:
        run.font.size = Pt(size)
    return paragraph


def _heading(doc: Document, text: str, level: int = 1) -> None:
    doc.add_heading(text, level=level)


def _bullets(doc: Document, items: list[str], *, numbered: bool = False) -> None:
    style = "List Number" if numbered else "List Bullet"
    for item in items:
        paragraph = doc.add_paragraph(style=style)
        paragraph.add_run(item)


def _table(
    doc: Document,
    headers: list[str],
    rows: list[list[str]],
    widths: list[float] | None = None,
    font_size: float = 9.2,
) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.autofit = False
    table.style = "Table Grid"
    for col, title in enumerate(headers):
        cell = table.rows[0].cells[col]
        cell.text = title
        _shade(cell, NAVY)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        _cell_margins(cell)
        for paragraph in cell.paragraphs:
            paragraph.paragraph_format.space_after = Pt(0)
            for run in paragraph.runs:
                run.font.name = "Calibri"
                run.font.size = Pt(font_size)
                run.font.bold = True
                run.font.color.rgb = _rgb(WHITE)
    _repeat_header(table.rows[0])
    _keep_row(table.rows[0])
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        _keep_row(table.rows[-1])
        for col, value in enumerate(values):
            cells[col].text = str(value)
            cells[col].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            _cell_margins(cells[col])
            _shade(cells[col], WHITE if row_index % 2 else PALE)
            for paragraph in cells[col].paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.03
                for run in paragraph.runs:
                    run.font.name = "Calibri"
                    run.font.size = Pt(font_size)
                    run.font.color.rgb = _rgb(INK)
                    if col == 0:
                        run.font.bold = True
                        run.font.color.rgb = _rgb(NAVY)
    if widths:
        for row in table.rows:
            for index, width in enumerate(widths):
                row.cells[index].width = Inches(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def _figure(doc: Document, filename: str, caption: str, width: float = 6.55) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(2)
    p.add_run().add_picture(str(DIAGRAMS / filename), width=Inches(width))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(7)
    run = cap.add_run(caption)
    run.font.name = "Calibri"
    run.font.size = Pt(8.5)
    run.font.italic = True
    run.font.color.rgb = _rgb(MUTED)


def _code(doc: Document, text: str) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    cell = table.cell(0, 0)
    _shade(cell, PALE)
    _cell_margins(cell, top=110, start=140, bottom=110, end=140)
    cell.text = ""
    p = cell.paragraphs[0]
    p.style = doc.styles["Report Code"]
    p.add_run(text)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def _callout(doc: Document, title: str, body: str) -> None:
    table = doc.add_table(rows=1, cols=1)
    cell = table.cell(0, 0)
    _shade(cell, LIGHT_BLUE)
    _cell_margins(cell, top=130, start=170, bottom=130, end=170)
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(title + "  ")
    r.bold = True
    r.font.color.rgb = _rgb(BLUE)
    p.add_run(body)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def cover(doc: Document) -> None:
    _paragraph(
        doc,
        "PROJECT DEFENCE REPORT",
        align=WD_ALIGN_PARAGRAPH.CENTER,
        color=GOLD,
        size=12,
        bold=True,
    )
    _paragraph(doc, "AeroReserve", align=WD_ALIGN_PARAGRAPH.CENTER, color=NAVY, size=31, bold=True)
    _paragraph(
        doc,
        "Flight Booking and Seat Allocation Platform",
        align=WD_ALIGN_PARAGRAPH.CENTER,
        color=BLUE,
        size=16,
    )
    _paragraph(
        doc,
        "Dated routes · 45-minute connections · cabin seats · one PNR per itinerary",
        align=WD_ALIGN_PARAGRAPH.CENTER,
        color=MUTED,
        size=10,
    )
    _paragraph(doc, "", size=5)
    _table(
        doc,
        ["Detail", "Entry"],
        [
            ["Course / subject", "Computational Thinking and Programming"],
            ["Project component", "Project Defence and Demonstration"],
            ["Submitted by", "____________________________________________"],
            ["Roll / registration number", "____________________________________________"],
            ["Department / institution", "____________________________________________"],
            ["Faculty guide", "____________________________________________"],
            ["Submission date", REPORT_DATE],
        ],
        widths=[2.05, 4.85],
        font_size=10,
    )
    _paragraph(doc, "", size=7)
    _paragraph(
        doc,
        "Python 3.11+  |  FastAPI  |  SQLAlchemy  |  SQLite default  |  Optional PostgreSQL",
        align=WD_ALIGN_PARAGRAPH.CENTER,
        color=MUTED,
        size=9,
    )
    doc.add_page_break()


def abstract(doc: Document) -> None:
    _heading(doc, "Abstract", 1)
    _paragraph(
        doc,
        "AeroReserve is a flight-booking academic demo developed by changing the Railway Reservation System into a flight-specific application. It searches a fictional dated timetable across seven Indian airports, supports Economy, Premium Economy, Business and First cabins, and returns direct or connecting itineraries with up to two connections. Every connection must use the same transfer airport and allow at least 45 minutes. One PNR represents all flight legs in an itinerary.",
    )
    _paragraph(
        doc,
        "The implementation consists of typed domain models, deterministic daily sample schedules, bounded route search, per-flight/cabin seat maps, three seat-allocation strategies, an itemised illustrative fare, SQLAlchemy persistence, a FastAPI REST API, a browser UI, a CLI and server-sent-event (SSE) live updates. SQLite creates itself on first run; PostgreSQL is configurable but was not integration-tested in this validation. The concurrency demonstration uses process-local locks and makes no multi-process correctness claim.",
    )
    _paragraph(
        doc,
        f"Validation on Python 3.13.16 completed with {TEST_COUNT} passing tests and {COVERAGE} combined statement/branch coverage. Ruff lint/format and strict mypy checks passed. A 50-thread, single-process race against the 36-seat Economy cabin confirmed 36 bookings, rejected 14 and assigned no duplicate seat labels. These results validate the tested demo configuration; they are not a production load test.",
    )
    _heading(doc, "Project at a glance", 2)
    _table(
        doc,
        ["Measure", "Verified project value"],
        [
            ["Sample schedule", "27 fictional daily flight templates · 7 airports"],
            ["Search scope", "1–3 legs · up to 2 connections · 45-minute minimum layover"],
            ["Cabins / seats per flight", "Economy 36 · Premium Economy 12 · Business 8 · First 4"],
            ["Booking model", "One 10-character PNR · ordered flight legs · cancellation"],
            [
                "Validation",
                f"{TEST_COUNT} tests · {COVERAGE} statement/branch coverage · ruff/mypy pass",
            ],
            ["Thread race", "50 contenders · 36 confirmed · 14 rejected · 0 duplicate seat labels"],
        ],
        widths=[2.05, 4.85],
        font_size=9.2,
    )
    _callout(
        doc,
        "Scope note",
        "All schedules and fare values are fictional. The app does not connect to an airline, sell tickets, process payments, or authenticate passengers.",
    )
    doc.add_page_break()


def contents(doc: Document) -> None:
    _heading(doc, "Contents", 1)
    entries = [
        "Abstract and project at a glance",
        "1. Introduction",
        "2. Requirements and scope",
        "3. Computational thinking and problem decomposition",
        "4. System design and data model",
        "5. Algorithms and complexity",
        "6. Implementation: API, browser and CLI",
        "7. Testing and quality engineering",
        "8. Deployment and reproducibility",
        "9. AI-assisted engineering record",
        "10. Validation results and interpretation",
        "11. User guide and demonstration plan",
        "12. Limitations and future work",
        "13. Conclusion",
        "References",
        "Appendix A. Project structure",
        "Appendix B. API reference and sample commands",
        "Appendix C. Formulae and validation record",
    ]
    for entry in entries:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(5)
        run = p.add_run(entry)
        run.font.name = "Calibri"
        run.font.size = Pt(11)
        run.font.color.rgb = _rgb(NAVY if entry[:1].isdigit() else INK)
    _paragraph(
        doc,
        "The contents are a section guide. Update the student and institution fields on the cover before submission.",
        italic=True,
        color=MUTED,
        size=9,
    )
    doc.add_page_break()


def chapter1(doc: Document) -> None:
    _heading(doc, "1. Introduction", 1)
    _heading(doc, "1.1 Background", 2)
    _paragraph(
        doc,
        "An airline booking application has to reason about dated flight instances rather than a rail journey’s route segments. A traveller can choose a direct service or a sequence of flights; each flight has its own seat inventory and a connecting itinerary must satisfy both airport continuity and a minimum transfer time. A booking must represent the complete trip rather than separate unrelated records.",
    )
    _heading(doc, "1.2 Problem statement", 2)
    _paragraph(
        doc,
        "Given two airports, a travel date, a cabin and a party size, the application must search its available sample schedule, reject impossible or too-short connections, estimate a fare, display cabin-specific availability, assign seats on each leg, persist a complete itinerary under one PNR, and release all leg seats on cancellation. Connected browser clients should observe committed booking changes without a manual refresh.",
    )
    _heading(doc, "1.3 Objectives", 2)
    _bullets(
        doc,
        [
            "Convert the original rail domain into airports, dated flights, cabins, flight legs, itineraries and PNR bookings.",
            "Provide repeatable fictional schedules and route/date search across direct and connecting trips.",
            "Enforce up to two connections, a 45-minute layover minimum and one PNR across one to three legs.",
            "Implement cabin-specific seat maps, explainable allocation policies, fare quotes, booking and cancellation.",
            "Deliver a working browser interface, REST endpoints, CLI, SSE updates and SQLite persistence without database setup.",
            "Test domain rules, APIs, live SSE and in-process threaded booking; document what the tests do not establish.",
        ],
    )
    _heading(doc, "1.4 Scope", 2)
    _table(
        doc,
        ["Included", "Not included"],
        [
            [
                "Seven fictional Indian airports and 27 daily flight templates.",
                "Live airline, airport, price or availability feeds.",
            ],
            [
                "Four cabins, dates, passenger count 1–6, itinerary search, booking/cancellation.",
                "Payments, e-ticket issuance, log-in, admin roles or identity checks.",
            ],
            [
                "Direct and 1-/2-connection routes; one PNR for all legs.",
                "International time zones, overnight connecting search or interline agreements.",
            ],
            [
                "Self-seeding SQLite; SQLAlchemy URL supports optional PostgreSQL.",
                "Production or multi-process seat consistency; PostgreSQL not integration-tested.",
            ],
        ],
        widths=[3.45, 3.45],
        font_size=9.1,
    )


def chapter2(doc: Document) -> None:
    _heading(doc, "2. Requirements and scope", 1)
    _heading(doc, "2.1 Functional requirements", 2)
    _table(
        doc,
        ["ID", "Requirement", "Implementation evidence"],
        [
            [
                "FR-01",
                "List sample airports and search by origin, destination and date.",
                "`/api/airports`; `/api/flights/search`; date limited to 180 days.",
            ],
            [
                "FR-02",
                "Offer cabin-specific inventory and seat maps.",
                "Four cabin enums and independent flight-instance seat labels.",
            ],
            [
                "FR-03",
                "Support direct and connected itineraries up to three legs.",
                "Bounded graph walk; maximum two connections.",
            ],
            [
                "FR-04",
                "Require at least 45 minutes at each transfer airport.",
                "Itinerary validator and route-search filter.",
            ],
            [
                "FR-05",
                "Use one PNR for all itinerary legs.",
                "One bookings row and ordered booking_legs rows per itinerary.",
            ],
            [
                "FR-06",
                "Quote, book, retrieve and cancel.",
                "Quote and booking API; PNR retrieval and cancellation routes.",
            ],
            [
                "FR-07",
                "Reflect committed changes in subscribed seat maps.",
                "Initial SSE snapshot plus booking/cancellation delta events.",
            ],
        ],
        widths=[0.65, 3.05, 3.2],
        font_size=8.7,
    )
    _heading(doc, "2.2 Quality requirements", 2)
    _bullets(
        doc,
        [
            "Zero-config local run: SQLite and automatic schema creation; no separate database is required for a demonstration.",
            "Maintain separation among domain/service, API, repository and self-contained browser UI.",
            "Validate input at API and service boundaries; use structured error responses and safe HTML escaping in the client.",
            "Keep the default schedule deterministic and demo data clearly identified as fictional.",
            "Run repeatable tests with temporary isolated databases; include lint, formatting and static typing.",
            "State the concurrency boundary plainly: only one application process is protected by its process-local locks.",
        ],
    )
    _heading(doc, "2.3 User roles", 2)
    _paragraph(
        doc,
        "The demo has a single unauthenticated browser user. There is no role model or real passenger-data security boundary. Test and demo runs should use fictional names only; access control, retention, audit governance and privacy review are future requirements before any real personal data is entered.",
    )


def chapter3(doc: Document) -> None:
    _heading(doc, "3. Computational thinking and problem decomposition", 1)
    _heading(doc, "3.1 Decomposition", 2)
    _paragraph(
        doc,
        "The complete request is separated into airport/date validation, schedule materialisation, itinerary search, connection-rule validation, inventory lookup, seat selection, fare calculation, persistence and event delivery. Each component has a direct test seam: route search accepts a date, seat selection accepts an occupied-label set, fare calculation accepts an itinerary and cabin, and the repository can be pointed at an isolated SQLite database.",
    )
    _heading(doc, "3.2 Pattern recognition and abstraction", 2)
    _paragraph(
        doc,
        "Repeated daily templates are converted into dated `Flight` values. A flight instance is the inventory boundary: `(flight_id, cabin)` selects its seat set. Itineraries are ordered tuples of one to three flights, while cabin layouts are reusable templates. This abstraction distinguishes two identical seat labels on different dates or flight numbers.",
    )
    _heading(doc, "3.3 Algorithmic thinking", 2)
    _paragraph(
        doc,
        "The search problem is a bounded route walk through a directed schedule graph. Every partial path carries its current airport and last arrival time. A connection is feasible only when the next departure is at least 45 minutes later; a visited-airport set prevents cycles. The allocation problem is a small per-flight assignment with a party size of at most six, so simple deterministic policies are suitable and easy to explain.",
    )
    _heading(doc, "3.4 Evaluation criteria", 2)
    _bullets(
        doc,
        [
            "Correct route: adjacent flights join at the same airport, each layover meets the minimum, and the final destination is reached.",
            "Inventory: every assigned label is valid for the selected cabin and not occupied by another confirmed booking on that flight instance.",
            "Atomic itinerary: a booking and all its legs commit together under one PNR.",
            "Consistency boundary: the test proves the current single-process thread path only; it is not a distributed-systems proof.",
            "Transparency: fare line items, algorithm rules, demo inputs and untested integrations are disclosed.",
        ],
    )


def chapter4(doc: Document) -> None:
    _heading(doc, "4. System design and data model", 1)
    _heading(doc, "4.1 Layered architecture", 2)
    _paragraph(
        doc,
        "The browser calls relative API paths, allowing the same HTML to run from the FastAPI process. FastAPI validates request payloads with Pydantic and delegates blocking service calls to its thread pool. The application service uses immutable domain values, the sample schedule, the cabin allocator, a process-local flight-lock manager and a SQLAlchemy repository. On success, it publishes bounded in-process events for SSE subscribers.",
    )
    _figure(
        doc,
        "architecture.png",
        "Figure 1. Current runtime components and the explicitly process-local lock/event paths.",
    )
    _heading(doc, "4.2 Persistence model", 2)
    _figure(
        doc,
        "data_model.png",
        "Figure 2. Persisted booking, leg and audit tables. Flight templates and cabin layouts are code constants.",
    )
    _paragraph(
        doc,
        "`bookings` stores the PNR, booking status, travel date, cabin, route, passenger payload, fare payload and timestamps. `booking_legs` stores one ordered row for every flight in the itinerary, including the flight ID, route, scheduled times, cabin, assigned seat labels and distance. `audit_log` records `BOOK` and `CANCEL` actions. Cancelled rows remain available for retrieval and auditing; only `CONFIRMED` booking legs consume seats.",
    )
    _heading(doc, "4.3 Booking lifecycle", 2)
    _figure(
        doc,
        "booking_flow.png",
        "Figure 3. Booking and cancellation flow. A PNR covers every leg; the DB transaction precedes SSE publication.",
    )


def chapter5(doc: Document) -> None:
    _heading(doc, "5. Algorithms and complexity", 1)
    _heading(doc, "5.1 Dated itinerary search", 2)
    _paragraph(
        doc,
        "For one requested date, the service builds an outgoing-flight map, then recursively walks candidate flights to depth at most three. It stops when it reaches the destination, rejects a previously visited airport, and advances the earliest allowed next departure by 45 minutes from the previous arrival. Results are sorted by arrival time, number of connections, total distance and flight-ID tuple.",
    )
    _figure(
        doc,
        "route_network.png",
        "Figure 4. Sample direct, one-stop and two-stop routes from VGA to DEL; schedules are fictional.",
    )
    _heading(doc, "5.2 Seat-allocation policies", 2)
    _table(
        doc,
        ["Strategy", "Rule", "Trade-off"],
        [
            [
                "First-fit",
                "Take the first free labels in row/column order.",
                "Predictable and simple; may split a party across aisle/rows.",
            ],
            [
                "Window-first",
                "Order window seats first, then row and column.",
                "Prefers a window; does not promise seats together.",
            ],
            [
                "Family-aware",
                "Find a contiguous same-row block on one side of aisle; otherwise fall back to first-fit.",
                "Keeps small parties together where possible; remains a greedy heuristic.",
            ],
        ],
        widths=[1.15, 3.0, 2.75],
        font_size=8.8,
    )
    _heading(doc, "5.3 Fare arithmetic", 2)
    _paragraph(
        doc,
        "For each leg, base amount is `distance_km × cabin_rate × passenger_count`. The project sums line amounts, adds 5% illustrative tax and adds ₹99 per passenger as a service fee. For one Economy passenger on the 1,370 km VGA–DEL sample direct flight: ₹5,206.00 base + ₹260.30 tax + ₹99.00 service fee = ₹5,565.30. The figures are constants for this software demonstration and are not real airline prices or tax advice.",
    )
    _heading(doc, "5.4 Complexity", 2)
    _table(
        doc,
        ["Operation", "Bound / cost", "Notes"],
        [
            ["Build dated schedule map", "O(E)", "E is the small number of daily scheduled edges."],
            [
                "Bounded route walk",
                "O(Σ bᵏ), k ≤ 3",
                "b is candidate branching; visited airports avoid cycles.",
            ],
            ["Sort matching paths", "O(P log P)", "P is the number of returned itineraries."],
            [
                "Seat selection",
                "O(C) family/first-fit; O(C log C) window sort",
                "C ≤ 36 seats per cabin in this model.",
            ],
            ["Quote / lock ordering", "O(L) quote; O(L log L) lock ordering", "L ≤ 3 flight legs."],
        ],
        widths=[1.7, 1.8, 3.4],
        font_size=8.7,
    )
    _paragraph(
        doc,
        "The graph and cabin sizes are intentionally small; the project makes no large-network performance claim. Search bounds derive from the maximum three legs, and not from a benchmark on airline-scale data.",
        italic=True,
        color=MUTED,
        size=9,
    )


def chapter6(doc: Document) -> None:
    _heading(doc, "6. Implementation: API, browser and CLI", 1)
    _heading(doc, "6.1 Domain and schedule", 2)
    _paragraph(
        doc,
        "The domain is represented by typed dataclasses and a string enum for cabin classes. `Itinerary` validates leg count, airport continuity and layover duration. `india_today()` uses Asia/Kolkata for date boundaries. The generated daily schedule has 27 fictional templates across VGA, HYD, BLR, BOM, DEL, MAA and GOI; all service labels are explicitly sample data.",
    )
    _heading(doc, "6.2 REST API", 2)
    _paragraph(
        doc,
        "FastAPI routes cover airport listing, search, cabin seat maps, quotes, booking, PNR retrieval, cancellation, health, audit and per-flight SSE. Pydantic models validate passenger name/age, party size, cabin and flight-leg count. Domain exceptions map to status-coded JSON responses. The generated OpenAPI interface is available at `/docs` while the local server is running [1, 2].",
    )
    _heading(doc, "6.3 Browser UI", 2)
    _paragraph(
        doc,
        "The self-contained page provides origin/destination/date, cabin and passenger controls; itinerary cards; individual flight-leg summaries; a seat-map preview; seat allocation policy selection; an itemised quote; and booking confirmation. The header API, Live, DB and Subs pills are working buttons that open live popovers rather than decorative status indicators. Native `EventSource` listens for the initial seat snapshot and for booking/cancellation events; the client refreshes the authoritative seat map on updates [6].",
    )
    _heading(doc, "6.4 CLI", 2)
    _paragraph(
        doc,
        "`aerobooking airports`, `search`, `book`, `booking`, `cancel` and `stress` provide a terminal path through the sample data. Search emits JSON. The `stress` command is an in-process threaded correctness demo; it is not a benchmark or a multi-process test.",
    )
    _code(
        doc,
        "aerobooking search --from VGA --to DEL --connections 2\naerobooking book --from VGA --to DEL -p 'Asha Rao:22'\naerobooking booking <PNR>\naerobooking cancel <PNR>",
    )


def chapter7(doc: Document) -> None:
    _heading(doc, "7. Testing and quality engineering", 1)
    _paragraph(
        doc,
        f"The current automated test suite reports {TEST_COUNT} passing tests in the local validation environment. It combines example-based route/API/booking tests, Hypothesis-generated seat/fare properties, an actual Uvicorn SSE integration test and a thread race. Branch measurement is enabled by coverage configuration.",
    )
    _table(
        doc,
        ["Area", "Evidence"],
        [
            [
                "Routing",
                "Direct, one-stop and two-stop paths; airport continuity; 45-minute lower bound; date and airport validation.",
            ],
            [
                "Seats and fares",
                "Cabin capacities, label uniqueness, occupied-seat exclusion, aisle-aware selection, fare decomposition.",
            ],
            [
                "Booking lifecycle",
                "One PNR for multiple legs, inventory by date/cabin, cancellation and re-bookable seats.",
            ],
            [
                "API/UI transport",
                "Search, quote, booking, validation errors, health/audit and root HTML; real SSE snapshot followed by a booking delta.",
            ],
            [
                "Property tests",
                "Hypothesis checks unique/free labels and fare monotonicity over generated inputs.",
            ],
            [
                "Concurrency",
                "50 worker threads race for the 36 Economy seats; database active assignments are checked for duplicates.",
            ],
        ],
        widths=[1.55, 5.35],
        font_size=8.9,
    )
    _heading(doc, "7.1 Observed quality checks", 2)
    _table(
        doc,
        ["Check", "Observed result"],
        [
            ["pytest", f"{TEST_COUNT} passed"],
            ["pytest-cov with branch measurement", COVERAGE],
            ["ruff check src tests", "All checks passed"],
            ["ruff format src tests", "Formatting clean"],
            ["mypy (strict)", "Success: no issues found in 9 source files"],
        ],
        widths=[2.7, 4.2],
        font_size=9,
    )
    _paragraph(
        doc,
        "Coverage is a code-execution metric, not a proof of correctness. PostgreSQL integration, Windows scripts, Docker Compose, CI on GitHub, multi-process locking and external live feeds were not executed in this validation.",
        italic=True,
        color=MUTED,
        size=9,
    )


def chapter8(doc: Document) -> None:
    _heading(doc, "8. Deployment and reproducibility", 1)
    _heading(doc, "8.1 Local demo", 2)
    _code(
        doc,
        'python -m venv .venv\n. .venv/bin/activate\npython -m pip install -e ".[dev]"\npython -m uvicorn aerobooking.api:app --reload --host 0.0.0.0 --port 8000\n# Open http://localhost:8000',
    )
    _paragraph(
        doc,
        "On Windows, equivalent `.bat` and PowerShell launchers are included. The commands have not been run on Windows in this validation environment. The server binds to all interfaces for local/demo access; open `http://localhost:8000` with the explicit port.",
    )
    _heading(doc, "8.2 Database configuration", 2)
    _paragraph(
        doc,
        "The default engine uses a self-seeding SQLite file. `AEROBOOKING_SQLITE_PATH` changes its location. `DATABASE_URL` takes precedence and accepts a SQLAlchemy PostgreSQL URL with Psycopg 3, for example `postgresql+psycopg://user:password@host:5432/aerobooking`. SQLAlchemy is the persistence toolkit; SQLite is the zero-setup route used by tests [3, 4, 7].",
    )
    _heading(doc, "8.3 Container and CI configuration", 2)
    _paragraph(
        doc,
        "A Dockerfile and Compose stack are provided. Compose configures a PostgreSQL service and waits for its health check before starting the API. GitHub Actions is configured for Python 3.11 and 3.12, lint, formatting, mypy, pytest/coverage and wheel build. These configuration files are included for reproducibility, but were not run against Docker or a remote CI runner during this report generation.",
    )
    _callout(
        doc,
        "Production boundary",
        "The current JSON seat payload and process-local RLocks do not prevent two independent application processes from selling the same seat. Keep the app at one worker for the demo; implement shared transactional inventory locking and add a multi-instance race test before scaling out.",
    )


def chapter9(doc: Document) -> None:
    _heading(doc, "9. AI-assisted engineering record", 1)
    _paragraph(
        doc,
        "AI assistance was used during specification refinement, implementation, test creation and document preparation. The required project scope was made explicit before coding: a flight-booking platform, dated/cabin-specific availability, connecting itineraries, a 45-minute minimum, one PNR across all legs, fictional schedules, zero-config SQLite and optional PostgreSQL. Generated work was treated as a draft; claims in the report are limited to code and validation evidence in this repository.",
    )
    _heading(doc, "9.1 Review and correction examples", 2)
    _bullets(
        doc,
        [
            "Aisle-aware family seat selection was given a dedicated test so a physical aisle is not counted as a contiguous seat gap; when no block exists, the implementation explicitly falls back to first-fit.",
            "The live API was exercised end-to-end: the SSE test asserts an initial 36-seat snapshot, posts a booking, then checks the booking event and 35-seat refreshed map.",
            "A code review found that refreshing results immediately after booking could clear the confirmation panel; the UI refresh now preserves the PNR result while clearing the prior itinerary selection.",
            "The concurrency statement was narrowed to one process after reviewing the lock manager and JSON seat persistence; no multi-worker correctness claim is made.",
            "PostgreSQL and Windows launchers are described as provided but unverified; no unsupported benchmark, airline feed or payment claim is added.",
        ],
    )
    _heading(doc, "9.2 Responsible-use note", 2)
    _paragraph(
        doc,
        "The student should replace placeholders, review every technical statement, run the demo on the submission machine and disclose AI assistance in the format required by the institution. This generated record describes the current artifact and does not assert that AI output was accepted without testing.",
    )


def chapter10(doc: Document) -> None:
    _heading(doc, "10. Validation results and interpretation", 1)
    _heading(doc, "10.1 Single-process booking race", 2)
    _paragraph(
        doc,
        "The stress command was run against a clean, isolated SQLite file for fictional flight AS110 in Economy. Fifty threads attempted one passenger each at the same time. The cabin has 36 seats. The observed result was:",
    )
    _table(
        doc,
        ["Contenders", "Capacity", "Confirmed", "Rejected", "Duplicate labels"],
        [["50 threads", "36", "36", "14", "0"]],
        widths=[1.55, 1.2, 1.25, 1.2, 1.7],
        font_size=9.5,
    )
    _paragraph(
        doc,
        "This is a correctness result for concurrent threads handled by one service process in the local demo database. It does not establish throughput, latency under load, crash recovery, multi-process safety, remote Postgres behaviour or production readiness. The zero-duplicate assertion is checked against active assignments read back from the database.",
    )
    _heading(doc, "10.2 Test and tool results", 2)
    _paragraph(
        doc,
        f"The suite completed with {TEST_COUNT} passing tests and {COVERAGE} branch-aware coverage. Ruff and mypy completed cleanly. The live SSE integration test used Uvicorn and an isolated SQLite database. No airline- or fare-provider connection exists; all search results are generated from in-repository fictional templates.",
    )
    _heading(doc, "10.3 Practical interpretation", 2)
    _paragraph(
        doc,
        "The implementation demonstrates a working end-to-end path: date and route search → itinerary selection → quote → cabin seat assignment for each flight → single-PNR persistence → API response → live seat update → cancellation. The tested envelope is deliberately narrow: a single-process academic demo with a small static network and single-aisle cabin constants.",
    )


def chapter11(doc: Document) -> None:
    _heading(doc, "11. User guide and demonstration plan", 1)
    _table(
        doc,
        ["Time", "Action", "What to show / say"],
        [
            [
                "0–1 min",
                "Start server and open `http://localhost:8000`.",
                "SQLite initializes automatically; point out the explicit port and fictional-schedule badge.",
            ],
            [
                "1–2 min",
                "Search VGA → DEL for tomorrow; choose Economy.",
                "Show the direct flight and alternatives. Explain cabin/date availability.",
            ],
            [
                "2–3 min",
                "Select a one-connection option, then a two-connection option.",
                "Point out each flight leg and the 45-minute minimum; one PNR will cover all legs.",
            ],
            [
                "3–4 min",
                "Change cabin and passenger count; preview seats and fare.",
                "Show seat-map state, clickable header popovers and one of the allocation strategies.",
            ],
            [
                "4–5 min",
                "Book the selected itinerary with fictional names.",
                "Read back the PNR and per-leg seats; explain one booking row and ordered leg rows.",
            ],
            [
                "5–6 min",
                "Open a second tab before the booking and observe the event.",
                "A second selected seat map receives the update through SSE, without a page refresh.",
            ],
            [
                "6–7 min",
                "Cancel via API/CLI or use another clean demo PNR.",
                "Seats are released on every leg; cancellation is persisted and published.",
            ],
            [
                "7–8 min",
                "Run tests / clean-database stress command.",
                "45 tests, 93.1% branch-aware coverage; 50 vs 36 race has zero duplicate labels, single process only.",
            ],
        ],
        widths=[0.7, 2.4, 3.8],
        font_size=8.3,
    )
    _heading(doc, "11.1 Useful commands", 2)
    _code(
        doc,
        "pytest -q\nruff check src tests && ruff format --check src tests && mypy\nAEROBOOKING_SQLITE_PATH=./stress-demo.db aerobooking stress --flight AS110 --threads 50\n# Open http://localhost:8000 and http://localhost:8000/docs",
    )
    _paragraph(
        doc,
        "Use fictional passenger names during a defence. A fresh stress database is necessary to reproduce 36 confirmations; previously booked seats reduce the number of successes. The command demonstrates one-process threads only.",
        italic=True,
        color=MUTED,
        size=9,
    )


def chapter12(doc: Document) -> None:
    _heading(doc, "12. Limitations and future work", 1)
    _table(
        doc,
        ["Current limitation", "Potential next step"],
        [
            [
                "Schedules are static, fictional templates repeated daily.",
                "Integrate a licensed timetable/inventory provider and persist schedule versions.",
            ],
            [
                "Cabin capacity/layout is generic code, not aircraft/seat configuration.",
                "Model aircraft configurations, seat attributes, blocked seats and inventory revisions.",
            ],
            [
                "Fare formula uses project constants; no repricing rules.",
                "Integrate a fare engine, currencies, discounts, fees and auditable tax sources.",
            ],
            [
                "No login, payment, e-ticketing or passenger-data controls.",
                "Add authentication, access policy, payment sandbox, data-minimisation and security review.",
            ],
            [
                "Process-local locks and event bus; JSON seat payload is not a shared uniqueness key.",
                "Use authoritative inventory rows / database locking, shared event delivery and multi-process race tests.",
            ],
            [
                "PostgreSQL URL path and Docker stack are untested here.",
                "Add a PostgreSQL service to integration CI and execute transaction/rollback/migration tests.",
            ],
            [
                "One-day, same-timezone route graph; no overnight/date-crossing connections.",
                "Search multi-day, timezone-aware itineraries with minimum connection rules per airport.",
            ],
        ],
        widths=[3.25, 3.65],
        font_size=8.5,
    )
    _paragraph(
        doc,
        "A production roadmap should begin with inventory correctness and identity/privacy requirements, not with adding more UI. Shared seat-level constraints and operational monitoring must be tested before multiple API workers or replicas are enabled.",
    )


def chapter13(doc: Document) -> None:
    _heading(doc, "13. Conclusion", 1)
    _paragraph(
        doc,
        "AeroReserve changes the original railway project into a flight-specific booking platform, rather than only renaming it. The demonstrated features include dated route search across a fictional sample network, cabin-specific availability, direct and connecting itineraries with a 45-minute minimum, one PNR across all legs, seat-allocation strategies, fare breakdown, cancellation, SQLite persistence, REST endpoints, a browser UI and live SSE refreshes.",
    )
    _paragraph(
        doc,
        f"The deliverable is supported by {TEST_COUNT} passing tests, {COVERAGE} branch-aware coverage, clean lint/type checks and a 50-thread single-process race result of 36 confirmed, 14 rejected and zero duplicate seat labels. The report also states the untested PostgreSQL/Windows/CI/container paths and the process-local concurrency limit. This is a working, reviewable academic demo and a clear base for future development—not a live airline reservation system.",
    )


def references(doc: Document) -> None:
    _heading(doc, "References", 1)
    refs = [
        "[1] FastAPI, ‘FastAPI Documentation,’ https://fastapi.tiangolo.com/ (accessed 13 October 2026).",
        "[2] Pydantic, ‘Pydantic Documentation,’ https://docs.pydantic.dev/latest/ (accessed 13 October 2026).",
        "[3] SQLAlchemy, ‘SQLAlchemy 2.0 Documentation,’ https://docs.sqlalchemy.org/en/20/ (accessed 13 October 2026).",
        "[4] SQLite Consortium, ‘SQLite Documentation,’ https://www.sqlite.org/docs.html (accessed 13 October 2026).",
        "[5] Python Software Foundation, ‘threading — Thread-based parallelism,’ https://docs.python.org/3/library/threading.html (accessed 13 October 2026).",
        "[6] WHATWG, ‘HTML Living Standard: Server-sent events,’ https://html.spec.whatwg.org/multipage/server-sent-events.html (accessed 13 October 2026).",
        "[7] Psycopg, ‘Psycopg 3 documentation,’ https://www.psycopg.org/psycopg3/docs/ (accessed 13 October 2026).",
        "[8] AeroReserve source code, test suite, and local CLI validation in this project package, October 2026.",
    ]
    for item in refs:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.24)
        p.paragraph_format.first_line_indent = Inches(-0.24)
        p.paragraph_format.space_after = Pt(6)
        p.add_run(item)


def appendices(doc: Document) -> None:
    _heading(doc, "Appendix A. Project structure", 1)
    _code(
        doc,
        "src/aerobooking/models.py       domain models and itinerary rules\nsrc/aerobooking/schedule.py     fictional dated sample flights\nsrc/aerobooking/availability.py cabin layouts and allocation policies\nsrc/aerobooking/db.py           SQLAlchemy schema and engine setup\nsrc/aerobooking/repository.py   booking, leg and audit persistence\nsrc/aerobooking/service.py      search, fares, booking, locks and SSE bus\nsrc/aerobooking/api.py          FastAPI endpoints and app factory\nsrc/aerobooking/cli.py          command-line demo\nsrc/aerobooking/web/index.html  self-contained browser interface\ntests/                          45 automated tests\ndocs/                           report, slide deck, defence pack and diagrams\nscripts/                        report/slide/diagram generators",
    )
    _heading(doc, "Appendix B. API reference and sample commands", 1)
    _table(
        doc,
        ["Method", "Route", "Purpose"],
        [
            ["GET", "/api/airports", "List fictional airports."],
            ["GET", "/api/flights/search", "Search itinerary options and fare/seat availability."],
            ["GET", "/api/flights/{flight_id}/seats", "Fetch a cabin seat map for a date."],
            ["GET", "/api/flights/{flight_id}/events", "SSE snapshot and live updates."],
            ["POST", "/api/bookings/quote", "Quote the selected flight legs."],
            ["POST", "/api/bookings", "Book 1–3 joined flight legs under one PNR."],
            ["GET", "/api/bookings/{pnr}", "Retrieve a PNR."],
            ["POST", "/api/bookings/{pnr}/cancel", "Cancel all legs under a PNR."],
            [
                "GET",
                "/api/health",
                "Database label, active bookings, subscribers and lock counters.",
            ],
            ["GET", "/api/admin/audit", "Retrieve recent booking/cancellation events."],
        ],
        widths=[0.7, 2.95, 3.25],
        font_size=8.7,
    )
    _code(
        doc,
        "curl 'http://localhost:8000/api/flights/search?origin=VGA&destination=DEL&travel_date=2026-10-14&cabin=Economy&max_connections=2&passengers=2'\n# Browser UI: http://localhost:8000\n# OpenAPI:    http://localhost:8000/docs",
    )
    _heading(doc, "Appendix C. Formulae and validation record", 1)
    _table(
        doc,
        ["Item", "Value / command"],
        [
            ["Itinerary maximum", "3 legs = 2 connections"],
            ["Minimum layover", "45 minutes"],
            ["Cabin capacities", "Economy 36 · Premium Economy 12 · Business 8 · First 4"],
            [
                "Fare",
                "Σ(distance × cabin rate × party size) + 5% illustrative tax + ₹99 × party size",
            ],
            ["Tests", "45 passed on Python 3.13.16"],
            ["Coverage", "93.1% statement/branch result with pytest-cov"],
            [
                "Stress",
                "50 threads · 36 confirmed · 14 rejected · 0 duplicate labels · one process",
            ],
            ["Static checks", "ruff clean; mypy strict clean"],
        ],
        widths=[1.75, 5.15],
        font_size=9,
    )
    _paragraph(
        doc,
        "Reproduce the recorded stress result with a fresh local database file and one application process. Do not cite it as evidence for multi-process, PostgreSQL, network-scale or airline-inventory behaviour.",
        italic=True,
        color=MUTED,
        size=9,
    )


def main() -> None:
    DOCS.mkdir(parents=True, exist_ok=True)
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.68)
    section.left_margin = Inches(0.65)
    section.right_margin = Inches(0.65)
    section.header_distance = Inches(0.36)
    section.footer_distance = Inches(0.35)
    _set_styles(doc)
    _header_footer(section)
    doc.core_properties.title = "AeroReserve Project Defence Report"
    doc.core_properties.subject = (
        "Flight booking, itinerary search and seat allocation academic demo"
    )
    doc.core_properties.author = "AeroReserve project team"
    doc.core_properties.keywords = "airline booking, PNR, itinerary, seat allocation, FastAPI"

    cover(doc)
    abstract(doc)
    contents(doc)
    for chapter in (
        chapter1,
        chapter2,
        chapter3,
        chapter4,
        chapter5,
        chapter6,
        chapter7,
        chapter8,
        chapter9,
        chapter10,
        chapter11,
        chapter12,
        chapter13,
        references,
        appendices,
    ):
        chapter(doc)
    doc.save(OUT)
    print(f"Wrote {OUT} ({OUT.stat().st_size:,} bytes)")


if __name__ == "__main__":
    main()
