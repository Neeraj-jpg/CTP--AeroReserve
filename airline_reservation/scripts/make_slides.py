"""Generate the 11-slide AeroReserve defence deck (PPTX)."""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"
DIAGRAMS = DOCS / "diagrams"
OUT = DOCS / "DEFENCE_SLIDES.pptx"
W, H = 13.333, 7.5

NAVY = "102B46"
BLUE = "1463A5"
CYAN = "14B8C5"
GOLD = "F4A62A"
INK = "172B3E"
MUTED = "536B80"
PALE = "F3F7FA"
LIGHT_BLUE = "E8F3F9"
LIGHT_CYAN = "E7F8F7"
LIGHT_GOLD = "FFF5E4"
WHITE = "FFFFFF"
GREEN = "16835D"
RED = "B74652"
FONT = "Aptos"
BLANK = None


def rgb(color: str) -> RGBColor:
    return RGBColor.from_string(color)


def add_rect(
    slide,
    x: float,
    y: float,
    w: float,
    h: float,
    fill: str,
    *,
    radius: bool = False,
    line: str | None = None,
):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    shape.line.color.rgb = rgb(line or fill)
    shape.line.width = Pt(0.8)
    return shape


def add_text(
    slide,
    x: float,
    y: float,
    w: float,
    h: float,
    text: str,
    *,
    size: float = 16,
    color: str = INK,
    bold: bool = False,
    align=PP_ALIGN.LEFT,
    valign=MSO_ANCHOR.TOP,
    margin: float = 0.06,
    font_name: str = FONT,
    italic: bool = False,
):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(margin)
    tf.margin_right = Inches(margin)
    tf.margin_top = Inches(margin)
    tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = valign
    paragraph = tf.paragraphs[0]
    paragraph.alignment = align
    paragraph.space_after = Pt(0)
    run = paragraph.add_run()
    run.text = text
    run.font.name = font_name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = rgb(color)
    return shape


def add_bullet(
    slide,
    x: float,
    y: float,
    w: float,
    text: str,
    *,
    size: float = 14,
    color: str = INK,
    bullet_color: str = CYAN,
):
    dot = slide.shapes.add_shape(
        MSO_SHAPE.OVAL, Inches(x), Inches(y + 0.12), Inches(0.09), Inches(0.09)
    )
    dot.fill.solid()
    dot.fill.fore_color.rgb = rgb(bullet_color)
    dot.line.color.rgb = rgb(bullet_color)
    return add_text(slide, x + 0.18, y, w - 0.18, 0.48, text, size=size, color=color, margin=0.01)


def add_line(
    slide,
    x1: float,
    y1: float,
    x2: float,
    y2: float,
    color: str = BLUE,
    width: float = 2.0,
    arrow: bool = False,
):
    shape = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2)
    )
    shape.line.color.rgb = rgb(color)
    shape.line.width = Pt(width)
    if arrow:
        shape.line.end_arrowhead = True
    return shape


def add_image(slide, filename: str, x: float, y: float, w: float, h: float | None = None):
    path = DIAGRAMS / filename
    if h is None:
        return slide.shapes.add_picture(str(path), Inches(x), Inches(y), width=Inches(w))
    return slide.shapes.add_picture(
        str(path), Inches(x), Inches(y), width=Inches(w), height=Inches(h)
    )


def new_slide(prs: Presentation):
    return prs.slides.add_slide(prs.slide_layouts[6])


def header(slide, section: str, title: str, number: int):
    add_rect(slide, 0, 0, W, 1.06, NAVY)
    add_rect(slide, 0, 1.06, W, 0.045, GOLD)
    add_text(
        slide,
        0.56,
        0.13,
        12.15,
        0.25,
        section.upper(),
        size=9.5,
        color=CYAN,
        bold=True,
        margin=0.01,
    )
    add_text(
        slide,
        0.56,
        0.39,
        12.15,
        0.55,
        title,
        size=26,
        color=WHITE,
        bold=True,
        margin=0.01,
        valign=MSO_ANCHOR.MIDDLE,
    )
    add_text(
        slide,
        0.58,
        7.12,
        8.6,
        0.18,
        "AeroReserve · Project Defence · Fictional sample platform",
        size=8.5,
        color=MUTED,
        margin=0.01,
    )
    add_text(
        slide,
        12.05,
        7.08,
        0.7,
        0.24,
        f"{number:02d} / 11",
        size=8.5,
        color=MUTED,
        align=PP_ALIGN.RIGHT,
        margin=0.01,
    )


def card(
    slide,
    x,
    y,
    w,
    h,
    title: str,
    body: str,
    *,
    fill: str = PALE,
    accent: str = CYAN,
    title_size: float = 16,
    body_size: float = 12.5,
):
    add_rect(slide, x, y, w, h, fill, radius=True, line="D6E2EB")
    add_rect(slide, x, y, 0.09, h, accent)
    add_text(
        slide,
        x + 0.22,
        y + 0.16,
        w - 0.42,
        0.42,
        title,
        size=title_size,
        color=NAVY,
        bold=True,
        margin=0.02,
    )
    add_text(
        slide,
        x + 0.22,
        y + 0.62,
        w - 0.42,
        h - 0.76,
        body,
        size=body_size,
        color=MUTED,
        margin=0.02,
    )


def pill(slide, x: float, y: float, w: float, label: str, fill: str, color: str = NAVY):
    add_rect(slide, x, y, w, 0.33, fill, radius=True, line=fill)
    add_text(
        slide,
        x + 0.05,
        y + 0.025,
        w - 0.1,
        0.26,
        label,
        size=9.5,
        color=color,
        bold=True,
        align=PP_ALIGN.CENTER,
        valign=MSO_ANCHOR.MIDDLE,
        margin=0.01,
    )


def title_slide(prs: Presentation):
    slide = new_slide(prs)
    add_rect(slide, 0, 0, W, H, NAVY)
    add_rect(slide, 0.62, 0.64, 0.11, 4.05, GOLD)
    add_text(
        slide,
        1.04,
        0.76,
        5.8,
        0.34,
        "PROJECT DEFENCE · 2026",
        size=12,
        color=CYAN,
        bold=True,
        margin=0.01,
    )
    add_text(
        slide, 1.02, 1.40, 7.0, 0.95, "AeroReserve", size=42, color=WHITE, bold=True, margin=0.01
    )
    add_text(
        slide,
        1.05,
        2.43,
        7.45,
        0.76,
        "Flight Booking and\nSeat Allocation Platform",
        size=25,
        color="D6E7F2",
        margin=0.01,
    )
    add_text(
        slide,
        1.06,
        3.55,
        6.9,
        0.44,
        "Dated routes  ·  45-minute connections  ·  one PNR per itinerary",
        size=14,
        color=WHITE,
        margin=0.01,
    )
    # Route motif.
    coords = [(8.45, 2.05, "VGA"), (10.10, 1.45, "HYD"), (11.85, 2.05, "DEL")]
    add_line(slide, 8.75, 2.0, 9.82, 1.55, CYAN, 3, True)
    add_line(slide, 10.40, 1.55, 11.56, 2.0, GOLD, 3, True)
    for x, y, code in coords:
        circle = slide.shapes.add_shape(
            MSO_SHAPE.OVAL, Inches(x - 0.34), Inches(y - 0.34), Inches(0.68), Inches(0.68)
        )
        circle.fill.solid()
        circle.fill.fore_color.rgb = rgb(LIGHT_BLUE)
        circle.line.color.rgb = rgb(CYAN)
        add_text(
            slide,
            x - 0.27,
            y - 0.12,
            0.54,
            0.24,
            code,
            size=10,
            color=NAVY,
            bold=True,
            align=PP_ALIGN.CENTER,
            valign=MSO_ANCHOR.MIDDLE,
            margin=0,
        )
    add_text(
        slide,
        8.1,
        2.68,
        4.1,
        0.35,
        "Fictional sample schedule · IST",
        size=11,
        color="B8CCDB",
        align=PP_ALIGN.CENTER,
        margin=0.01,
    )
    add_rect(slide, 1.03, 4.55, 11.28, 1.25, "173B59", radius=True, line="315674")
    add_text(
        slide,
        1.30,
        4.78,
        10.8,
        0.55,
        "Student: ____________________     Roll No.: ____________________\nDepartment / Institution: __________________________________________",
        size=13,
        color=WHITE,
        margin=0.02,
    )
    add_text(
        slide,
        1.06,
        6.82,
        6.4,
        0.25,
        "Computational Thinking and Programming",
        size=10.5,
        color="B8CCDB",
        margin=0.01,
    )
    add_text(
        slide,
        11.0,
        6.82,
        1.3,
        0.25,
        "13 OCT 2026",
        size=9.5,
        color=CYAN,
        align=PP_ALIGN.RIGHT,
        margin=0.01,
    )


def slide_problem(prs: Presentation):
    slide = new_slide(prs)
    header(slide, "01 · Problem", "Flight booking needs flight-specific rules", 2)
    card(
        slide,
        0.62,
        1.45,
        5.45,
        2.15,
        "Not just a renamed rail UI",
        "A flight itinerary has dated instances, separate cabin inventory on each leg, airport transfers and flight schedules. One PNR must represent the complete trip.",
        fill=LIGHT_BLUE,
        accent=BLUE,
        title_size=17,
        body_size=14,
    )
    card(
        slide,
        6.32,
        1.45,
        6.36,
        2.15,
        "Core constraints",
        "Origin and destination · travel date · cabin\nDirect or up to 2 connections (3 legs)\nAt least 45 minutes at every transfer\nOne PNR and leg-specific seats",
        fill=LIGHT_CYAN,
        accent=CYAN,
        title_size=17,
        body_size=14,
    )
    # Three decisions.
    card(
        slide,
        0.62,
        3.93,
        3.82,
        1.6,
        "01 · Search",
        "Fictional daily schedule\n7 airports · 27 flight templates",
        fill=PALE,
        accent=CYAN,
        title_size=14,
        body_size=12,
    )
    card(
        slide,
        4.75,
        3.93,
        3.82,
        1.6,
        "02 · Connect",
        "Same transfer airport\n≥45-minute minimum layover",
        fill=LIGHT_GOLD,
        accent=GOLD,
        title_size=14,
        body_size=12,
    )
    card(
        slide,
        8.88,
        3.93,
        3.82,
        1.6,
        "03 · Reserve",
        "Cabin seats per flight leg\nOne PNR for the itinerary",
        fill=LIGHT_BLUE,
        accent=BLUE,
        title_size=14,
        body_size=12,
    )
    add_rect(slide, 0.62, 5.9, 12.06, 0.58, NAVY, radius=True, line=NAVY)
    add_text(
        slide,
        0.86,
        6.04,
        11.55,
        0.27,
        "Scope boundary: no live airline feed · no payment processing · no passenger login",
        size=12,
        color=WHITE,
        bold=True,
        align=PP_ALIGN.CENTER,
        valign=MSO_ANCHOR.MIDDLE,
        margin=0.01,
    )


def slide_workflow(prs: Presentation):
    slide = new_slide(prs)
    header(slide, "02 · Requirements", "A complete passenger journey", 3)
    steps = [
        (0.62, "1", "Search", "Choose airports, travel date, cabin and party size."),
        (3.78, "2", "Compare", "Review direct and connecting options, times, seats and fare."),
        (6.94, "3", "Book", "Allocate cabin seats separately on each leg under one PNR."),
        (10.10, "4", "Manage", "Retrieve or cancel; live seat maps refresh on updates."),
    ]
    for x, n, title, body in steps:
        add_rect(slide, x, 1.7, 2.62, 2.35, WHITE, radius=True, line="D6E2EB")
        badge = slide.shapes.add_shape(
            MSO_SHAPE.OVAL, Inches(x + 0.2), Inches(1.94), Inches(0.48), Inches(0.48)
        )
        badge.fill.solid()
        badge.fill.fore_color.rgb = rgb(CYAN)
        badge.line.color.rgb = rgb(CYAN)
        add_text(
            slide,
            x + 0.28,
            2.05,
            0.32,
            0.2,
            n,
            size=12,
            bold=True,
            color=NAVY,
            align=PP_ALIGN.CENTER,
            margin=0,
        )
        add_text(
            slide, x + 0.2, 2.60, 2.2, 0.36, title, size=18, bold=True, color=NAVY, margin=0.01
        )
        add_text(slide, x + 0.2, 3.08, 2.2, 0.7, body, size=12.2, color=MUTED, margin=0.01)
    # cabin row
    add_text(
        slide,
        0.66,
        4.43,
        2.2,
        0.33,
        "Cabin inventory per sample flight",
        size=13,
        color=NAVY,
        bold=True,
        margin=0.01,
    )
    cabins = [
        ("Economy", "36", LIGHT_BLUE),
        ("Premium Economy", "12", LIGHT_CYAN),
        ("Business", "8", LIGHT_GOLD),
        ("First", "4", "F5EDF6"),
    ]
    x = 0.66
    for name, count, fill in cabins:
        width = 2.82
        add_rect(slide, x, 4.9, width, 0.85, fill, radius=True, line="D6E2EB")
        add_text(
            slide,
            x + 0.13,
            5.02,
            width - 0.26,
            0.28,
            name,
            size=11.2,
            color=NAVY,
            bold=True,
            margin=0.01,
        )
        add_text(
            slide,
            x + 0.13,
            5.34,
            width - 0.26,
            0.25,
            f"{count} seats",
            size=12,
            color=BLUE,
            bold=True,
            margin=0.01,
        )
        x += width + 0.17
    add_text(
        slide,
        0.66,
        6.12,
        11.95,
        0.35,
        "Passengers: 1–6 · dates: today through the next 180 days (IST) · schedules/fare constants are fictional.",
        size=10.6,
        color=MUTED,
        italic=True,
        margin=0.01,
    )


def slide_architecture(prs: Presentation):
    slide = new_slide(prs)
    header(slide, "03 · Design", "Layered design with explicit boundaries", 4)
    add_image(slide, "architecture.png", 1.94, 1.17, 9.45)
    add_text(
        slide,
        1.02,
        6.48,
        11.3,
        0.34,
        "Locks and event queues are process-local. PostgreSQL URL support is provided, but multi-process consistency is not claimed.",
        size=10.5,
        color=RED,
        bold=True,
        align=PP_ALIGN.CENTER,
        margin=0.01,
    )


def slide_data(prs: Presentation):
    slide = new_slide(prs)
    header(slide, "04 · Data model", "One booking row, ordered rows for each leg", 5)
    add_image(slide, "data_model.png", 1.63, 1.18, 10.05)
    add_text(
        slide,
        1.02,
        6.48,
        11.3,
        0.30,
        "Cabin seat occupancy is derived from confirmed booking legs for a dated flight and cabin.",
        size=10.5,
        color=MUTED,
        align=PP_ALIGN.CENTER,
        margin=0.01,
    )


def slide_search(prs: Presentation):
    slide = new_slide(prs)
    header(slide, "05 · Algorithms", "Bounded route search with a transfer-time rule", 6)
    add_image(slide, "route_network.png", 0.55, 1.3, 7.55)
    add_rect(slide, 8.33, 1.48, 4.35, 4.85, NAVY, radius=True, line=NAVY)
    add_text(
        slide,
        8.62,
        1.78,
        3.75,
        0.45,
        "Search sequence",
        size=19,
        color=WHITE,
        bold=True,
        margin=0.01,
    )
    bullets = [
        "Build outgoing-flight map for the travel date.",
        "Walk candidate paths to a maximum of 3 legs.",
        "Reject airport cycles and broken airport joins.",
        "Require next departure ≥ previous arrival + 45 min.",
        "Sort by arrival time, stops, distance and flight IDs.",
    ]
    y = 2.38
    for item in bullets:
        add_bullet(slide, 8.62, y, 3.72, item, size=11.6, color="E4EEF5", bullet_color=CYAN)
        y += 0.72
    add_text(
        slide,
        0.72,
        6.25,
        7.2,
        0.38,
        "Small fictional network; depth is capped at three legs. No large-network performance claim.",
        size=10.4,
        color=MUTED,
        italic=True,
        margin=0.01,
    )


def slide_seats_fare(prs: Presentation):
    slide = new_slide(prs)
    header(slide, "06 · Pricing and seats", "Cabin-specific inventory, explainable policies", 7)
    add_rect(slide, 0.64, 1.45, 5.78, 2.68, LIGHT_BLUE, radius=True, line="D6E2EB")
    add_text(
        slide,
        0.92,
        1.68,
        5.15,
        0.38,
        "Allocation strategies",
        size=18,
        color=NAVY,
        bold=True,
        margin=0.01,
    )
    add_bullet(slide, 0.95, 2.22, 5.0, "First-fit: earliest free labels in row order.", size=12.4)
    add_bullet(slide, 0.95, 2.82, 5.0, "Window-first: window seats, then row/column.", size=12.4)
    add_bullet(
        slide,
        0.95,
        3.42,
        5.0,
        "Family-aware: same-row block off the aisle; otherwise first-fit.",
        size=12.4,
    )
    add_rect(slide, 6.77, 1.45, 5.92, 2.68, LIGHT_GOLD, radius=True, line="E9D9B6")
    add_text(
        slide,
        7.05,
        1.68,
        5.35,
        0.38,
        "Illustrative Economy quote · 1 passenger",
        size=17,
        color=NAVY,
        bold=True,
        margin=0.01,
    )
    fare_rows = [
        ("VGA–DEL base · 1,370 km × ₹3.80", "₹5,206.00"),
        ("Illustrative tax · 5%", "₹260.30"),
        ("Service fee · ₹99 / passenger", "₹99.00"),
    ]
    y = 2.25
    for label, value in fare_rows:
        add_text(slide, 7.05, y, 3.85, 0.32, label, size=11.1, color=MUTED, margin=0.01)
        add_text(
            slide,
            10.95,
            y,
            1.35,
            0.32,
            value,
            size=11.2,
            color=INK,
            bold=True,
            align=PP_ALIGN.RIGHT,
            margin=0.01,
        )
        y += 0.48
    add_rect(slide, 7.04, 3.63, 5.28, 0.34, NAVY, radius=True, line=NAVY)
    add_text(
        slide,
        7.17,
        3.69,
        5.0,
        0.22,
        "TOTAL  ₹5,565.30  ·  sample formula only",
        size=11,
        color=WHITE,
        bold=True,
        align=PP_ALIGN.CENTER,
        margin=0.01,
    )
    add_text(
        slide,
        0.7,
        4.58,
        4.0,
        0.33,
        "Seat counts per dated sample flight",
        size=13.5,
        color=NAVY,
        bold=True,
        margin=0.01,
    )
    capacities = [("Economy", "36"), ("Premium Economy", "12"), ("Business", "8"), ("First", "4")]
    x = 0.72
    for name, count in capacities:
        add_rect(slide, x, 5.05, 2.85, 0.72, WHITE, radius=True, line="D6E2EB")
        add_text(
            slide, x + 0.15, 5.13, 1.78, 0.24, name, size=10.4, color=MUTED, bold=True, margin=0.01
        )
        add_text(
            slide,
            x + 2.0,
            5.14,
            0.55,
            0.25,
            count,
            size=13,
            color=BLUE,
            bold=True,
            align=PP_ALIGN.RIGHT,
            margin=0.01,
        )
        x += 3.04
    add_text(
        slide,
        0.72,
        6.14,
        11.9,
        0.34,
        "A connecting booking receives one seat assignment per leg. Fare rates and tax are demo constants, not real fares/tax advice.",
        size=10.2,
        color=MUTED,
        italic=True,
        margin=0.01,
    )


def slide_lifecycle(prs: Presentation):
    slide = new_slide(prs)
    header(slide, "07 · Booking and live updates", "Validate → allocate → commit → publish", 8)
    add_image(slide, "booking_flow.png", 0.91, 1.38, 11.5)
    add_rect(slide, 1.08, 5.78, 5.3, 0.62, LIGHT_CYAN, radius=True, line="C9E8E5")
    add_text(
        slide,
        1.28,
        5.93,
        4.9,
        0.30,
        "SSE snapshot + booking/cancellation event per flight leg",
        size=11.3,
        color=NAVY,
        bold=True,
        align=PP_ALIGN.CENTER,
        margin=0.01,
    )
    add_rect(slide, 6.8, 5.78, 5.3, 0.62, LIGHT_GOLD, radius=True, line="E9D9B6")
    add_text(
        slide,
        7.0,
        5.93,
        4.9,
        0.30,
        "One PNR and one DB transaction cover the itinerary",
        size=11.3,
        color=NAVY,
        bold=True,
        align=PP_ALIGN.CENTER,
        margin=0.01,
    )


def slide_results(prs: Presentation):
    slide = new_slide(prs)
    header(slide, "08 · Verification", "Test evidence, not production claims", 9)
    stats = [
        (0.68, "45", "tests passed", BLUE),
        (3.8, "93.1%", "branch-aware coverage", CYAN),
        (7.2, "ruff", "clean", GREEN),
        (10.25, "mypy", "strict · clean", GOLD),
    ]
    for x, big, small, color in stats:
        add_rect(slide, x, 1.52, 2.45, 1.12, PALE, radius=True, line="D6E2EB")
        add_text(
            slide,
            x + 0.12,
            1.68,
            2.2,
            0.46,
            big,
            size=25,
            color=color,
            bold=True,
            align=PP_ALIGN.CENTER,
            margin=0.01,
        )
        add_text(
            slide,
            x + 0.12,
            2.18,
            2.2,
            0.24,
            small,
            size=10.4,
            color=MUTED,
            align=PP_ALIGN.CENTER,
            margin=0.01,
        )
    add_text(
        slide,
        0.72,
        3.03,
        6.5,
        0.38,
        "Single-process seat race · 50 simultaneous threads",
        size=16,
        color=NAVY,
        bold=True,
        margin=0.01,
    )
    add_rect(slide, 0.76, 3.66, 8.1, 0.54, "E7EDF2", radius=True, line="E7EDF2")
    add_rect(slide, 0.76, 3.66, 5.832, 0.54, GREEN, radius=True, line=GREEN)
    # cover the rounded join to make the ratio bar squared at the inside boundary
    add_rect(slide, 6.52, 3.66, 0.08, 0.54, GREEN)
    add_text(
        slide,
        0.94,
        3.77,
        5.35,
        0.25,
        "36 confirmed · all Economy seats",
        size=11.3,
        color=WHITE,
        bold=True,
        align=PP_ALIGN.CENTER,
        margin=0.01,
    )
    add_text(
        slide,
        6.68,
        3.77,
        2.0,
        0.25,
        "14 rejected",
        size=10.8,
        color=NAVY,
        bold=True,
        align=PP_ALIGN.CENTER,
        margin=0.01,
    )
    add_rect(slide, 9.28, 3.3, 3.38, 1.65, LIGHT_CYAN, radius=True, line="C9E8E5")
    add_text(
        slide,
        9.52,
        3.57,
        2.9,
        0.56,
        "0",
        size=33,
        color=GREEN,
        bold=True,
        align=PP_ALIGN.CENTER,
        margin=0.01,
    )
    add_text(
        slide,
        9.5,
        4.19,
        2.92,
        0.3,
        "duplicate seat labels",
        size=12,
        color=NAVY,
        bold=True,
        align=PP_ALIGN.CENTER,
        margin=0.01,
    )
    add_bullet(
        slide,
        0.78,
        4.77,
        8.0,
        "End-to-end SSE snapshot then booking delta tested over Uvicorn.",
        size=11.5,
    )
    add_bullet(
        slide,
        0.78,
        5.37,
        8.0,
        "Search, layover, cabin, quote, PNR and cancellation tests; Hypothesis properties.",
        size=11.5,
    )
    add_rect(slide, 0.74, 6.12, 11.94, 0.48, LIGHT_GOLD, radius=True, line="E9D9B6")
    add_text(
        slide,
        0.95,
        6.23,
        11.5,
        0.23,
        "Thread correctness in one process only. No multi-worker or PostgreSQL integration result.",
        size=11,
        color=NAVY,
        bold=True,
        align=PP_ALIGN.CENTER,
        margin=0.01,
    )


def slide_limits(prs: Presentation):
    slide = new_slide(prs)
    header(slide, "09 · Boundaries", "What this demo does — and does not — claim", 10)
    card(
        slide,
        0.68,
        1.48,
        5.75,
        1.72,
        "Working today",
        "Route/date search · cabin seat maps · booking/cancellation\nOne PNR for 1–3 legs · live SSE updates\nZero-config SQLite demo",
        fill=LIGHT_CYAN,
        accent=CYAN,
        title_size=16,
        body_size=12.3,
    )
    card(
        slide,
        6.88,
        1.48,
        5.75,
        1.72,
        "Explicitly not built",
        "Live airline data · real pricing\nPayment and ticket issue · passenger login\nMulti-worker shared inventory guarantee",
        fill=LIGHT_GOLD,
        accent=GOLD,
        title_size=16,
        body_size=12.3,
    )
    add_text(
        slide,
        0.75,
        3.62,
        4.0,
        0.35,
        "Next engineering steps",
        size=16,
        color=NAVY,
        bold=True,
        margin=0.01,
    )
    roadmap = [
        (0.72, "1", "Shared seat rows", "DB uniqueness / transaction locks"),
        (3.88, "2", "Integration tests", "Postgres + multiple workers"),
        (7.04, "3", "External services", "Licensed schedule / fare source"),
        (10.20, "4", "Security", "Auth, privacy and payments review"),
    ]
    for x, n, title, body in roadmap:
        add_rect(slide, x, 4.22, 2.4, 1.48, PALE, radius=True, line="D6E2EB")
        pill(slide, x + 0.15, 4.4, 0.38, n, LIGHT_BLUE, BLUE)
        add_text(
            slide, x + 0.15, 4.88, 2.1, 0.28, title, size=12.2, color=NAVY, bold=True, margin=0.01
        )
        add_text(slide, x + 0.15, 5.23, 2.1, 0.38, body, size=10, color=MUTED, margin=0.01)
    add_text(
        slide,
        0.78,
        6.14,
        11.9,
        0.33,
        "AI-assisted work was tested/reviewed; PostgreSQL, Docker, remote CI and Windows launchers remain unverified.",
        size=10.4,
        color=MUTED,
        italic=True,
        margin=0.01,
    )


def slide_demo(prs: Presentation):
    slide = new_slide(prs)
    header(slide, "10 · Demonstration", "Run the demo · questions", 11)
    add_rect(slide, 0.68, 1.48, 7.45, 3.05, NAVY, radius=True, line=NAVY)
    add_text(
        slide, 0.98, 1.78, 6.8, 0.34, "Start locally", size=18, color=CYAN, bold=True, margin=0.01
    )
    command = 'python -m pip install -e ".[dev]"\npython -m uvicorn aerobooking.api:app --reload --host 0.0.0.0 --port 8000'
    add_text(
        slide,
        0.98,
        2.32,
        6.75,
        0.82,
        command,
        size=10.5,
        color=WHITE,
        font_name="Consolas",
        margin=0.02,
    )
    add_rect(slide, 1.0, 3.47, 6.76, 0.62, CYAN, radius=True, line=CYAN)
    add_text(
        slide,
        1.16,
        3.62,
        6.42,
        0.28,
        "Open http://localhost:8000",
        size=15,
        color=NAVY,
        bold=True,
        align=PP_ALIGN.CENTER,
        margin=0.01,
    )
    add_text(
        slide,
        0.98,
        4.22,
        6.74,
        0.23,
        "API documentation: http://localhost:8000/docs",
        size=10.8,
        color="C8D9E6",
        margin=0.01,
    )
    add_rect(slide, 8.54, 1.48, 4.08, 3.05, LIGHT_BLUE, radius=True, line="D6E2EB")
    add_text(
        slide,
        8.84,
        1.78,
        3.5,
        0.34,
        "Eight-minute walkthrough",
        size=17,
        color=NAVY,
        bold=True,
        margin=0.01,
    )
    actions = [
        "1 · Search VGA → DEL",
        "2 · Compare 0/1/2 stops",
        "3 · Preview seats + fare",
        "4 · Book one PNR",
        "5 · Show second-tab SSE",
        "6 · Cancel and rerun tests",
    ]
    y = 2.32
    for action in actions:
        add_text(slide, 8.9, y, 3.38, 0.28, action, size=11.5, color=INK, margin=0.01)
        y += 0.34
    add_text(
        slide,
        1.02,
        5.35,
        11.25,
        0.55,
        "Thank you",
        size=28,
        color=NAVY,
        bold=True,
        align=PP_ALIGN.CENTER,
        margin=0.01,
    )
    add_text(
        slide,
        1.02,
        5.96,
        11.25,
        0.38,
        "Questions?",
        size=20,
        color=BLUE,
        bold=True,
        align=PP_ALIGN.CENTER,
        margin=0.01,
    )
    add_text(
        slide,
        1.02,
        6.57,
        11.25,
        0.24,
        "AeroReserve is a fictional academic demo, not a live airline booking service.",
        size=9.5,
        color=MUTED,
        italic=True,
        align=PP_ALIGN.CENTER,
        margin=0.01,
    )


def main() -> None:
    prs = Presentation()
    prs.slide_width = Inches(W)
    prs.slide_height = Inches(H)
    prs.core_properties.title = "AeroReserve Project Defence"
    prs.core_properties.subject = "Flight Booking and Seat Allocation Platform"
    prs.core_properties.author = "AeroReserve project team"
    title_slide(prs)
    for build in (
        slide_problem,
        slide_workflow,
        slide_architecture,
        slide_data,
        slide_search,
        slide_seats_fare,
        slide_lifecycle,
        slide_results,
        slide_limits,
        slide_demo,
    ):
        build(prs)
    prs.save(OUT)
    print(f"Wrote {OUT} ({OUT.stat().st_size:,} bytes, {len(prs.slides)} slides)")


if __name__ == "__main__":
    main()
